# -*- coding: utf-8 -*-
"""Check-in ROPRE em .pptx — o deck da reunião, na ordem do template da V4.

  python3 carregar/deck.py --checkin checkin.json --saida checkin.pptx

Sai um .pptx que o Google Slides abre e converte. A ordem dos blocos é a do
template: capa, índice, 01 Resultados, 02 Objetivos, 03 Premissas e Riscos,
04 Entregas, 05 Próximos Passos — e uma página final de fontes, que é o que
sustenta o número quando o cliente pergunta de onde veio.

Precisa de python-pptx (`python3 -m pip install --user python-pptx`).
"""
import argparse, json, os, sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from carregar.formato import moeda, moeda_curta, numero, percentual, variacao, valor, mes_curto  # noqa: E402

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

VERMELHO = RGBColor(0xE5, 0x09, 0x14)
PRETO = RGBColor(0x11, 0x11, 0x11)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
CINZA = RGBColor(0x70, 0x70, 0x70)
CINZA_CLARO = RGBColor(0xF2, 0xF2, 0xF2)
VERDE = RGBColor(0x1B, 0x7F, 0x3B)

L, T, W = Inches(0.7), Inches(0.55), Inches(11.9)
BLOCOS = [("01", "Resultados"), ("02", "Objetivos"), ("03", "Premissas e Riscos"),
          ("04", "Entregas"), ("05", "Próximos Passos")]


def _texto(slide, texto, esquerda, topo, largura, altura, tamanho=14, cor=PRETO,
           negrito=False, alinhamento=PP_ALIGN.LEFT, espaco=1.15):
    cx = slide.shapes.add_textbox(esquerda, topo, largura, altura)
    tf = cx.text_frame
    tf.word_wrap = True
    linhas = texto.split("\n") if isinstance(texto, str) else list(texto)
    for i, linha in enumerate(linhas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = str(linha)
        p.alignment = alinhamento
        p.line_spacing = espaco
        for run in p.runs:
            run.font.size = Pt(tamanho)
            run.font.color.rgb = cor
            run.font.bold = negrito
            run.font.name = "Arial"
    return cx


def _fundo(slide, cor):
    fundo = slide.background.fill
    fundo.solid()
    fundo.fore_color.rgb = cor


def _titulo(slide, titulo, subtitulo=None):
    _texto(slide, titulo, L, T, W, Inches(0.6), tamanho=26, negrito=True)
    if subtitulo:
        _texto(slide, subtitulo, L, T + Inches(0.62), W, Inches(0.4), tamanho=12, cor=CINZA)
    barra = slide.shapes.add_shape(1, L, T + Inches(1.05), Inches(1.1), Pt(4))
    barra.fill.solid()
    barra.fill.fore_color.rgb = VERMELHO
    barra.line.fill.background()
    barra.shadow.inherit = False


def _tabela(slide, cabecalho, linhas, topo, larguras=None, tamanho=11, altura_linha=0.32):
    n_col, n_lin = len(cabecalho), len(linhas) + 1
    larguras = larguras or [W / n_col] * n_col
    forma = slide.shapes.add_table(n_lin, n_col, L, topo, W, Inches(altura_linha * n_lin))
    tabela = forma.table
    for j, largura in enumerate(larguras):
        tabela.columns[j].width = Emu(int(largura))
    for j, titulo in enumerate(cabecalho):
        cel = tabela.cell(0, j)
        cel.text = str(titulo)
        cel.fill.solid()
        cel.fill.fore_color.rgb = PRETO
        cel.vertical_anchor = MSO_ANCHOR.MIDDLE
        for p in cel.text_frame.paragraphs:
            for run in p.runs:
                run.font.size = Pt(tamanho)
                run.font.bold = True
                run.font.color.rgb = BRANCO
                run.font.name = "Arial"
    for i, linha in enumerate(linhas, start=1):
        for j, celula in enumerate(linha):
            cel = tabela.cell(i, j)
            cel.text = "" if celula is None else str(celula)
            cel.fill.solid()
            cel.fill.fore_color.rgb = BRANCO if i % 2 else CINZA_CLARO
            cel.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cel.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(tamanho)
                    run.font.color.rgb = PRETO
                    run.font.name = "Arial"
    return tabela


def _cartao(slide, esquerda, topo, largura, rotulo, valor_txt, apoio=None, cor_apoio=CINZA):
    caixa = slide.shapes.add_shape(1, esquerda, topo, largura, Inches(1.55))
    caixa.fill.solid()
    caixa.fill.fore_color.rgb = CINZA_CLARO
    caixa.line.fill.background()
    caixa.shadow.inherit = False
    _texto(slide, rotulo.upper(), esquerda + Inches(0.2), topo + Inches(0.15), largura - Inches(0.4),
           Inches(0.3), tamanho=10, cor=CINZA, negrito=True)
    _texto(slide, valor_txt, esquerda + Inches(0.2), topo + Inches(0.45), largura - Inches(0.4),
           Inches(0.6), tamanho=22, negrito=True)
    if apoio:
        _texto(slide, apoio, esquerda + Inches(0.2), topo + Inches(1.08), largura - Inches(0.4),
               Inches(0.3), tamanho=11, cor=cor_apoio)


def _nova(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _divisoria(prs, numero_bloco, nome):
    s = _nova(prs)
    _fundo(s, PRETO)
    _texto(s, numero_bloco, L, Inches(2.4), Inches(3), Inches(1.2), tamanho=54, cor=VERMELHO, negrito=True)
    _texto(s, nome, L, Inches(3.5), W, Inches(1.0), tamanho=36, cor=BRANCO, negrito=True)
    trilha = " · ".join(f"{n} {t}" for n, t in BLOCOS)
    _texto(s, trilha, L, Inches(6.6), W, Inches(0.4), tamanho=11, cor=CINZA)
    return s


# ------------------------------------------------------------------ blocos

def _capa(prs, c):
    s = _nova(prs)
    _fundo(s, PRETO)
    p = c["periodo"]
    _texto(s, "R  O  P  R  E", L, Inches(1.5), W, Inches(0.5), tamanho=16, cor=VERMELHO, negrito=True)
    _texto(s, f"Check-in {p['cadencia']}", L, Inches(2.3), W, Inches(1.0), tamanho=44, cor=BRANCO, negrito=True)
    _texto(s, c["cliente"], L, Inches(3.4), W, Inches(0.8), tamanho=30, cor=BRANCO)
    rodape = f"{p['rotulo']}" + (f" · parcial até {p['ate']}" if p["parcial"] else "")
    _texto(s, rodape, L, Inches(4.4), W, Inches(0.5), tamanho=14, cor=CINZA)
    _texto(s, f"gerado em {c['gerado_em']} · fonte: {c.get('fonte')}", L, Inches(6.6), W, Inches(0.4),
           tamanho=10, cor=CINZA)


def _indice(prs, c):
    s = _nova(prs)
    _titulo(s, "ROPRE", "a estrutura do check-in")
    detalhe = {
        "Resultados": "período contra anterior · novo e recorrente · canal · atribuição",
        "Objetivos": "avanço das OKRs · step do projeto · o que foi acordado nas calls",
        "Premissas e Riscos": "o que assumimos · causa, risco e efeito · P × I",
        "Entregas": "realizadas · previstas · pendências abertas",
        "Próximos Passos": "o que vem, com dono e prazo",
    }
    topo = Inches(1.7)
    for n, nome in BLOCOS:
        _texto(s, n, L, topo, Inches(0.7), Inches(0.4), tamanho=18, cor=VERMELHO, negrito=True)
        _texto(s, nome, L + Inches(0.8), topo, Inches(3.4), Inches(0.4), tamanho=18, negrito=True)
        _texto(s, detalhe[nome], L + Inches(4.3), topo + Inches(0.05), Inches(7.2), Inches(0.4),
               tamanho=12, cor=CINZA)
        topo += Inches(0.95)


def _resultados(prs, c):
    r, p = c["resultados"], c["periodo"]
    ant, var = r.get("anterior") or {}, r.get("variacao") or {}

    s = _nova(prs)
    _titulo(s, "Resultados do período", f"{p['rotulo']} · comparado com {ant.get('periodo', {}).get('rotulo', '—')}")
    largura = Inches(2.78)
    for i, (rotulo, txt, apoio) in enumerate([
        ("Receita", moeda_curta(r["receita"]), f"{variacao(var.get('receita'))} vs anterior"),
        ("Vendas", numero(r["vendas"]), f"{variacao(var.get('vendas'))} vs anterior"),
        ("Ticket médio", moeda_curta(r["ticket"]), f"{variacao(var.get('ticket'))} vs anterior"),
        ("ROAS", numero(r["roas"], 2) if r.get("roas") else "não medido",
         f"{variacao(var.get('roas'))} vs anterior" if r.get("roas") else "base de mídia incompleta"),
    ]):
        _cartao(s, L + i * (largura + Inches(0.26)), Inches(1.5), largura, rotulo, txt, apoio)

    if r.get("atingimento") is not None:
        cor = VERDE if r["atingimento"] >= 1 else VERMELHO
        _texto(s, f"{percentual(r['atingimento'])} da meta do período", L, Inches(3.35), Inches(6), Inches(0.5),
               tamanho=20, negrito=True, cor=cor)
        _texto(s, f"meta {moeda(r['meta_periodo'])} · break-even {moeda(r['break_even_mes'])}/mês · "
                  f"resultado do período {moeda(r.get('resultado'))}",
               L, Inches(3.85), W, Inches(0.4), tamanho=12, cor=CINZA)

    linhas = [
        ["Novos", numero(r["novos"]["n"]), moeda(r["novos"]["receita"]), moeda(r["novos"]["ticket"])],
        ["Recorrentes", numero(r["recorrentes"]["n"]), moeda(r["recorrentes"]["receita"]),
         moeda(r["recorrentes"]["ticket"])],
        ["Total", numero(r["vendas"]), moeda(r["receita"]), moeda(r["ticket"])],
    ]
    _tabela(s, ["", "Vendas", "Receita", "Ticket"], linhas, Inches(4.5),
            larguras=[W * 0.34, W * 0.18, W * 0.24, W * 0.24])
    if r.get("participacao_recorrente") is not None:
        _texto(s, f"Recorrência é {percentual(r['participacao_recorrente'])} das vendas do período.",
               L, Inches(6.1), W, Inches(0.4), tamanho=12, cor=CINZA)

    # mês a mês
    if r.get("serie_mensal"):
        s = _nova(prs)
        _titulo(s, "Mês a mês", "faturamento por data de fechamento, já com a regra de atribuição do projeto")
        linhas = [[mes_curto(m), numero(d["vendas"]), moeda(d["receita"]), numero(d["novos"]),
                   numero(d["recorrentes"])] for m, d in r["serie_mensal"].items()]
        _tabela(s, ["Mês", "Vendas", "Receita", "Novos", "Recorrentes"], linhas, Inches(1.6),
                larguras=[W * 0.16, W * 0.16, W * 0.28, W * 0.2, W * 0.2])

    # ponte com a planilha
    if r.get("ponte_planilha"):
        s = _nova(prs)
        _titulo(s, "Planilha do cliente contra CRM",
                "a planilha lança só a venda nova; a recompra fica em linha separada e some do total")
        linhas = [[mes_curto(l["mes"]),
                   f"{numero(l['planilha_novas'])} · {moeda(l['planilha_receita_novas'], 0)}",
                   f"{numero(l['planilha_recorrentes'])} · {moeda(l['planilha_receita_recorrentes'], 0)}",
                   f"{numero(l['crm_novas'])} · {moeda(l['crm_receita_novas'], 0)}",
                   f"{numero(l['crm_recorrentes'])} · {moeda(l['crm_receita_recorrentes'], 0)}",
                   moeda(l["diferenca_receita"], 0)] for l in r["ponte_planilha"]]
        _tabela(s, ["Mês", "Planilha · novas", "Planilha · recorrentes", "CRM · novas", "CRM · recorrentes",
                    "Diferença"], linhas, Inches(1.7), tamanho=10)

    # safra / eficiência
    if r.get("safra"):
        s = _nova(prs)
        _titulo(s, "Eficiência da mídia", "por safra de criação — é a leitura que isola o mês que a mídia gerou")
        linhas = [[mes_curto(m), numero(d["criadas"]), percentual(d["maturidade"]), numero(d["ganhas"]),
                   percentual(d["taxa_ganho"], 1), moeda(d["investimento"], 0),
                   numero(d["eficiencia"], 1) if d["eficiencia"] else "não medido"]
                  for m, d in r["safra"].items()]
        _tabela(s, ["Safra", "Criadas", "Resolvidas", "Ganhas", "Taxa de ganho", "Investimento", "Eficiência"],
                linhas, Inches(1.7), tamanho=10)

    # canal + atribuição
    s = _nova(prs)
    _titulo(s, "Canal e atribuição", "de onde vem cada venda, e como ela está marcada no CRM")
    if r.get("por_canal"):
        linhas = [[canal, numero(d["vendas"]), moeda(d["receita"], 0), moeda(d["investimento"], 0),
                   numero(d["roas"], 2) if d["roas"] is not None else "—"]
                  for canal, d in r["por_canal"].items()]
        _tabela(s, ["Canal", "Vendas", "Receita", "Investimento", "ROAS"], linhas, Inches(1.6),
                larguras=[W * 0.3, W * 0.15, W * 0.2, W * 0.2, W * 0.15], tamanho=11)
    if r.get("cobertura_atribuicao"):
        topo = Inches(1.7 + 0.34 * (len(r.get("por_canal") or {}) + 2))
        _texto(s, f"Regra: {(c.get('regra_atribuicao') or {}).get('resumo', 'não declarada')}", L, topo,
               W, Inches(0.4), tamanho=11, cor=CINZA)
        linhas = [[marca, numero(n)] for marca, n in r["cobertura_atribuicao"].items()]
        _tabela(s, ["Marca no CRM", "Vendas"], linhas, topo + Inches(0.45),
                larguras=[W * 0.7, W * 0.3], tamanho=11)


def _objetivos(prs, c):
    o = c["objetivos"]
    s = _nova(prs)
    _titulo(s, "Objetivos", o.get("step") and f"Step do projeto: {o['step']}" or None)
    topo = Inches(1.6)
    if o.get("objetivo_smart"):
        _texto(s, o["objetivo_smart"], L, topo, W, Inches(1.0), tamanho=13)
        topo += Inches(1.1)
    if o.get("okrs"):
        linhas = [[kr["icone"], kr["kr"], valor(kr["meta"], kr["formato"]), valor(kr["realizado"], kr["formato"])]
                  for kr in o["okrs"]]
        _tabela(s, ["", "Key result", "Meta", "Realizado"], linhas, topo,
                larguras=[W * 0.06, W * 0.54, W * 0.2, W * 0.2], tamanho=11, altura_linha=0.42)
    if o.get("acordado_nas_calls"):
        s = _nova(prs)
        _titulo(s, "O que foi acordado com o cliente", "extraído das calls do período")
        _texto(s, [f"• {a}" for a in o["acordado_nas_calls"]], L, Inches(1.7), W, Inches(4.5), tamanho=14)


def _premissas_riscos(prs, c):
    pr = c["premissas_riscos"]
    s = _nova(prs)
    _titulo(s, "Premissas", "o que este check-in assume como verdade")
    _tabela(s, ["Premissa", "Valor", "Origem"],
            [[p["premissa"], p["valor"], p.get("origem", "—")] for p in pr["premissas"]],
            Inches(1.6), larguras=[W * 0.3, W * 0.3, W * 0.4], tamanho=11)

    s = _nova(prs)
    _titulo(s, "Riscos", "causa, risco e efeito — ordenados por probabilidade × impacto")
    linhas = []
    for risco in pr["riscos"][:8]:
        pi = f"{risco['probabilidade']} × {risco['impacto']} = {risco['score']}" if risco.get("score") else "—"
        linhas.append([risco.get("causa", "—"), risco.get("risco", "—"), risco.get("efeito") or "—", pi])
    _tabela(s, ["Causa", "Risco", "Efeito", "P × I"], linhas, Inches(1.6),
            larguras=[W * 0.28, W * 0.3, W * 0.28, W * 0.14], tamanho=10, altura_linha=0.55)


def _entregas(prs, c):
    e = c["entregas"]
    s = _nova(prs)
    _titulo(s, "Entregas realizadas", f"{len(e.get('realizadas') or [])} no período"
            + (f" · {e['horas']} dedicadas" if e.get("horas") else ""))
    itens = [f"✅ {i if isinstance(i, str) else i.get('entrega')}" for i in (e.get("realizadas") or [])]
    _texto(s, itens or ["Nenhuma entrega registrada no período."], L, Inches(1.6), W, Inches(5), tamanho=13)

    if e.get("previstas") or e.get("pendencias"):
        s = _nova(prs)
        _titulo(s, "Previstas e pendências", "o que está em andamento e o que trava")
        prev = [f"⏳ {i if isinstance(i, str) else i.get('entrega')}" for i in (e.get("previstas") or [])]
        _texto(s, ["PREVISTAS"] + (prev or ["—"]), L, Inches(1.6), W * 0.48, Inches(4.5), tamanho=13)
        pend = [f"• {p['dia']} · {p['pendencia']}" for p in (e.get("pendencias") or [])]
        _texto(s, ["PENDÊNCIAS ABERTAS"] + (pend or ["—"]), L + W * 0.52, Inches(1.6), W * 0.48,
               Inches(4.5), tamanho=13)


def _proximos(prs, c):
    s = _nova(prs)
    _titulo(s, "Próximos passos", "o que vem, com dono e prazo")
    itens = []
    for passo in (c.get("proximos_passos") or []):
        if isinstance(passo, str):
            itens.append(f"• {passo}")
        else:
            extra = " · ".join(x for x in [passo.get("dono"), passo.get("prazo"), passo.get("origem")] if x)
            itens.append(f"• {passo.get('passo')}" + (f"  ({extra})" if extra else ""))
    _texto(s, itens or ["Nenhum próximo passo registrado."], L, Inches(1.6), W, Inches(5), tamanho=14)


def _fontes(prs, c):
    s = _nova(prs)
    _titulo(s, "Fontes e o que não foi medido", "para o número ser defendido linha por linha")
    linhas = [
        "Faturamento lê data de fechamento no CRM; safra e eficiência leem data de criação.",
        "Recompra é o funil de recorrência do CRM, não o campo de venda base.",
        f"Extração: {c.get('fonte')} em {c.get('extraido_em')}.",
    ]
    _texto(s, linhas, L, Inches(1.6), W, Inches(1.4), tamanho=13)
    avisos = [f"• {a}" for a in (c.get("avisos") or [])] or ["• Sem lacunas de medição no período."]
    _texto(s, avisos, L, Inches(3.0), W, Inches(3.6), tamanho=12, cor=CINZA)


def montar(c, caminho):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    _capa(prs, c)
    _indice(prs, c)
    _divisoria(prs, "01", "Resultados")
    _resultados(prs, c)
    _divisoria(prs, "02", "Objetivos")
    _objetivos(prs, c)
    _divisoria(prs, "03", "Premissas e Riscos")
    _premissas_riscos(prs, c)
    _divisoria(prs, "04", "Entregas")
    _entregas(prs, c)
    _divisoria(prs, "05", "Próximos Passos")
    _proximos(prs, c)
    _fontes(prs, c)
    prs.save(caminho)
    return len(prs.slides._sldIdLst)


def main():
    ap = argparse.ArgumentParser(description="Renderiza o check-in ROPRE em .pptx")
    ap.add_argument("--checkin", required=True)
    ap.add_argument("--saida", required=True)
    a = ap.parse_args()
    with open(a.checkin, encoding="utf-8") as f:
        checkin = json.load(f)
    n = montar(checkin, a.saida)
    print(f"{a.saida}: {n} slides.")


if __name__ == "__main__":
    main()
